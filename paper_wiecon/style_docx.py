"""Post-process the pandoc docx into IEEE conference formatting:
letter page, two-column body, Times New Roman, styled headings/captions/
tables/references, centered display equations."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

TNR = "Times New Roman"

doc = Document("main_raw.docx")

# ---------- remove the stray "{26}" paragraph from thebibliography
for p in list(doc.paragraphs):
    if p.text.strip() == "26":
        p._element.getparent().remove(p._element)
        break

# ---------- fix doubled "Fig. Fig." / "Table Table" from resolved refs
for p in doc.paragraphs:
    if "Fig. Fig." in p.text or "Table Table" in p.text:
        for run in p.runs:
            run.text = (run.text.replace("Fig. Fig.", "Fig.")
                                 .replace("Table Table", "Table "))

# ---------- page setup
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.625)
sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1.0)

# ---------- base styles
normal = doc.styles["Normal"]
normal.font.name = TNR
normal.font.size = Pt(10)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)
rpr = normal.element.get_or_add_rPr()
rfonts = rpr.get_or_add_rFonts()
rfonts.set(qn("w:eastAsia"), TNR)

title = doc.styles["Title"]
title.font.name = TNR
title.font.size = Pt(24)
title.font.bold = False
title.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), TNR)

for hname, size, center, italic, smallcaps, before in [
        ("Heading 1", 10, True, False, True, 12),
        ("Heading 2", 10, False, True, False, 6)]:
    st = doc.styles[hname]
    st.font.name = TNR
    st.font.size = Pt(size)
    st.font.bold = False
    st.font.italic = italic
    pf = st.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(before)
    pf.space_after = Pt(4)
    r = st.element.get_or_add_rPr()
    rf = r.get_or_add_rFonts()
    if smallcaps:
        sc = OxmlElement("w:smallCaps")
        r.append(sc)

for cap in ["Table Caption", "Image Caption"]:
    st = doc.styles[cap]
    st.font.name = TNR
    st.font.size = Pt(8)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(6)

# ---------- title + two-column body section
title_p = doc.paragraphs[0]
assert title_p.style.name == "Title"
title_p.paragraph_format.space_after = Pt(6)

# continuous section break right after the title: first section = 1 column
p = title_p._element
sectPr = OxmlElement("w:sectPr")
pg = OxmlElement("w:pgSz")
pg.set(qn("w:w"), "12240")
pg.set(qn("w:h"), "15840")
sectPr.append(pg)
mg = OxmlElement("w:pgMar")
for k, v in [("top", "1080"), ("bottom", "1440"), ("left", "900"),
             ("right", "900"), ("header", "480"), ("footer", "480"),
             ("gutter", "0")]:
    mg.set(qn("w:" + k), v)
sectPr.append(mg)
ctype = OxmlElement("w:type")
ctype.set(qn("w:val"), "continuous")
sectPr.append(ctype)
p.addnext(sectPr)

# body: two columns on the final section
xpath = sec._sectPr.xpath("./w:cols")
if xpath:
    cols = xpath[0]
else:
    cols = OxmlElement("w:cols")
    sec._sectPr.append(cols)
cols.set(qn("w:num"), "2")
cols.set(qn("w:space"), "360")

# ---------- paragraph-level styling
in_refs = False
for para in doc.paragraphs:
    style = para.style.name
    txt = para.text.strip()

    if style == "Title":
        continue
    if style == "Heading 1" and txt == "Acknowledgment":
        in_refs_at_ack = True
    if style == "Heading 1" and txt == "References":
        in_refs = True

    if style in ("First Paragraph", "Body Text"):
        if txt.startswith("Abstract"):
            for run in para.runs:
                run.font.size = Pt(9)
                run.font.bold = True
            para.paragraph_format.space_before = Pt(6)
        elif txt.startswith("Index Terms"):
            for run in para.runs:
                run.font.size = Pt(9)
                run.font.bold = True
        if in_refs:
            for run in para.runs:
                run.font.size = Pt(8)
            para.paragraph_format.left_indent = Inches(0.18)
            para.paragraph_format.first_line_indent = Inches(-0.18)

    # display-equation paragraphs -> centered
    if para._element.xpath(".//m:oMathPara"):
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # figure paragraphs -> centered
    if style == "Captioned Figure":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ---------- table styling: IEEE rules (top/header/bottom only), 8 pt
def set_border(el, edge, sz=6):
    el.append(OxmlElement("w:" + edge))

for tbl in doc.tables:
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, sz in [("top", "8"), ("bottom", "8")]:
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), sz)
        e.set(qn("w:color"), "000000")
        borders.append(e)
    for edge in ["left", "right", "insideH", "insideV"]:
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "none")
        borders.append(e)
    tblPr.append(borders)
    for ri, row in enumerate(tbl.rows):
        if ri == 0:
            trPr = row._tr.get_or_add_trPr()
            tb = OxmlElement("w:tblBorders")
            for row_borders in []:
                pass
        for cell in row.cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.space_before = Pt(1)
                for run in para.runs:
                    run.font.name = TNR
                    run.font.size = Pt(8)

# header underline for each table: bottom border on first-row cells
for tbl in doc.tables:
    for cell in tbl.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tb = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "000000")
        tb.append(e)
        tcPr.append(tb)

doc.save("main.docx")
print("saved main.docx")
